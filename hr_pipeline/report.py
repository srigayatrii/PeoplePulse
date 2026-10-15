"""Excel report writer (openpyxl formatting on top of pandas)."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .transform import Metrics

NAVY = "1F3C6E"
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Arial")
BODY_FONT = Font(name="Arial", size=10)
THIN = Side(style="thin", color="D0D7E2")
SEVERITY_FILL = {"HIGH": "F8D7DA", "ERROR": "F8D7DA", "MEDIUM": "FFF3CD", "WARNING": "FFF3CD", "LOW": "E2E3E5", "REVIEW": "FFF3CD"}


def _style_sheet(ws, df: pd.DataFrame) -> None:
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.font = BODY_FONT
            cell.border = Border(bottom=THIN)
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
            if isinstance(cell.value, str) and cell.value in SEVERITY_FILL:
                cell.fill = PatternFill("solid", fgColor=SEVERITY_FILL[cell.value])
    for i, col in enumerate(df.columns, start=1):
        width = max([len(str(col))] + [len(str(v)) for v in df[col].head(200)]) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(max(width, 10), 60)
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def write_report(metrics: Metrics, dq_log: pd.DataFrame, validation: pd.DataFrame,
                 out_dir: Path | str, period: str) -> Path:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / f"people_report_{period.replace('-', '')}.xlsx"

    summary = pd.DataFrame({"metric": list(metrics.summary), "value": list(metrics.summary.values())})
    sheets = {
        "Summary": summary,
        "Headcount": metrics.headcount,
        "Turnover": metrics.turnover,
        "Payroll_Variance": metrics.payroll_variance,
        "Compliance_Alerts": metrics.compliance_gaps,
        "Validation": validation,
        "Data_Quality_Log": dq_log,
    }
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df = df.copy()
            for col in df.select_dtypes(include=["datetimetz", "datetime64[ns]"]).columns:
                df[col] = df[col].dt.date
            df.to_excel(xw, sheet_name=name, index=False)
            _style_sheet(xw.sheets[name], df)
    return path
